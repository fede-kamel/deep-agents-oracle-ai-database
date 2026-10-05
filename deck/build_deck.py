"""Build the deck on the Oracle FY26 corporate template.

The template (masters, layouts, theme) comes from an Oracle FY26 deck passed as
--template; every existing slide is dropped and the slides below are built on
its "Title/Cover_5" and "Light - Blank" layouts with the same type system: an
Oracle Sans kicker in Oracle red, a 30 pt title, a short red rule, and the
copyright footer. Run metrics are read from the event logs, so the numbers on
the slides are the numbers in the evidence.

    uv run --with python-pptx python deck/build_deck.py \
        --template ~/Desktop/litellm/LiteLLM-OCI-PPO-FY26.pptx \
        --out deck/Deep-Agents-Oracle-AI-Database-FY26.pptx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
FIG = ROOT / "docs" / "figures"
LOGOS = ROOT / "docs" / "figures-src" / "logos"
EVIDENCE = ROOT / "evidence"

RED = RGBColor(0xC7, 0x46, 0x34)
INK = RGBColor(0x2A, 0x2F, 0x2F)
BODY = RGBColor(0x57, 0x5C, 0x5C)
MUTED = RGBColor(0x8A, 0x8F, 0x8F)
TEAL = RGBColor(0x04, 0x53, 0x6F)
SAND = RGBColor(0xF4, 0xF2, 0xF0)
SAND_LINE = RGBColor(0xE2, 0xDF, 0xDB)
CARD_LINE = RGBColor(0xE6, 0xE2, 0xDE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x5F, 0x7D, 0x4F)
OCHRE = RGBColor(0xB0, 0x7D, 0x1F)
VIOLET = RGBColor(0x6B, 0x4F, 0xA0)
YELLOW = RGBColor(0xF0, 0xCC, 0x71)
FONT = "Oracle Sans"
MONO = "Consolas"
QUESTIONS = [
    "What has changed in kidney function, and which medications and monitoring deserve review?",
    "Why does she keep being readmitted, and what should the clinic address first?",
    "Which medications add to fall risk or duplicate each other?",
]
DOCTOR = {"executed": "approved", "needs_physician": "waiting", "rejected": "rejected"}
FOOTER = "Copyright © 2026, Oracle and/or its affiliates   —   Personal open-source work with synthetic data"


# ------------------------------------------------------------------ helpers --

def drop_all_slides(prs: Presentation) -> None:
    ids = prs.slides._sldIdLst
    for sid in list(ids):
        prs.part.drop_rel(sid.rId)
        ids.remove(sid)


def layout(prs, name):
    return next(lay for lay in prs.slide_layouts if lay.name == name)


def text(slide, x, y, w, h, runs, size=14, color=BODY, bold=False, align=PP_ALIGN.LEFT, font=FONT,
         anchor=MSO_ANCHOR.TOP, spacing=None, line=1.12):
    """runs: str, or list of paragraphs; a paragraph is a str or a list of (text, overrides) tuples."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    paragraphs = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line
        parts = [(para, {})] if isinstance(para, str) else para
        for chunk, o in parts:
            r = p.add_run()
            r.text = chunk
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", False)
            f.color.rgb = o.get("color", color)
            if spacing or o.get("spacing"):
                r._r.get_or_add_rPr().set("spc", str(o.get("spacing", spacing)))
        if isinstance(para, list) and para and para[0][1].get("space_after") is not None:
            p.space_after = Pt(para[0][1]["space_after"])
    return box


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def header(slide, kicker, title):
    text(slide, 0.72, 0.46, 11.9, 0.3, kicker.upper(), size=11.5, color=RED, bold=True, spacing=180)
    text(slide, 0.72, 0.78, 11.9, 1.0, title, size=30, color=INK, bold=True)
    rect(slide, 0.72, 1.40, 1.15, 0.06, RED)
    text(slide, 0.72, 7.06, 11.9, 0.3, FOOTER, size=8, color=MUTED)


def bullets(slide, x, y, w, h, items, size=14.5):
    paras = []
    for lead, rest in items:
        paras.append([("▪  ", {"color": RED, "bold": True, "size": size}), (lead + " ", {"bold": True, "color": INK, "size": size}),
                      (rest, {"color": BODY, "size": size})])
    box = text(slide, x, y, w, h, paras, size=size)
    for p in box.text_frame.paragraphs:
        p.space_after = Pt(11)
    return box


def card(slide, x, y, w, h, title, body, accent=RED, fill=SAND, title_size=15, body_size=12.5):
    rect(slide, x, y, w, h, fill, line=CARD_LINE)
    rect(slide, x, y, w, 0.07, accent)
    text(slide, x + 0.2, y + 0.2, w - 0.4, h - 0.3,
         [[(title, {"bold": True, "color": INK, "size": title_size})], [(body, {"color": BODY, "size": body_size})]], line=1.15)


def figure(slide, name, x, y, w, caption=None, frame=True):
    path = FIG / name
    if frame:
        rect(slide, x - 0.08, y - 0.08, w + 0.16, 0.1, WHITE)  # placeholder, resized below
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(w))
    if frame:
        border = slide.shapes[-2]
        border.top, border.left = pic.top - Inches(0.08), pic.left - Inches(0.08)
        border.width, border.height = pic.width + Inches(0.16), pic.height + Inches(0.16)
        border.line.color.rgb = CARD_LINE
        border.line.width = Pt(0.75)
    if caption:
        cy = Emu(pic.top + pic.height).inches + 0.14
        text(slide, x, cy, w, 0.5, caption, size=10.5, color=BODY, line=1.1)
        slide.shapes[-1].text_frame.paragraphs[0].runs[0].font.italic = True
    return pic


def code(slide, x, y, w, h, source, size=11.5):
    rect(slide, x, y, w, h, SAND, line=SAND_LINE)
    rect(slide, x, y, 0.07, h, TEAL)
    lines = source.strip("\n").splitlines()
    paras = []
    for ln in lines:
        stripped = ln.lstrip()
        color = MUTED if stripped.startswith("#") else (TEAL if stripped.split(" ")[0] in ("from", "import", "def", "return", "CREATE", "BEGIN", "END;", "SELECT", "FROM", "WHERE") else INK)
        paras.append([(ln or " ", {"font": MONO, "size": size, "color": color})])
    text(slide, x + 0.25, y + 0.2, w - 0.4, h - 0.3, paras, line=1.05)


def table(slide, x, y, w, rows, col_w, size=12, header_fill=TEAL):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.4 * len(rows)))
    tbl = shape.table
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if i == 0 else (WHITE if i % 2 else SAND)
            tf = cell.text_frame
            tf.paragraphs[0].text = ""
            r = tf.paragraphs[0].add_run()
            r.text = str(val)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.bold = i == 0 or j == 0
            r.font.color.rgb = WHITE if i == 0 else (INK if j == 0 else BODY)
            cell.margin_left = cell.margin_right = Inches(0.12)
    return tbl


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body


def repo(slide, paths):
    """Where to find it: the files behind the slide, so every claim is code you can open."""
    text(slide, 0.72, 6.72, 11.9, 0.3, [[("In the repo:  ", {"bold": True, "color": RED, "size": 10.5}),
                                         (paths, {"font": MONO, "color": BODY, "size": 10})]])


def steps(slide, x, y, w, h, items, gap=0.32, title_size=14, body_size=11.5):
    """A left-to-right flow of numbered boxes: items are (title, body, accent)."""
    n = len(items)
    bw = (w - gap * (n - 1)) / n
    for i, (title, body, accent) in enumerate(items):
        bx = x + i * (bw + gap)
        rect(slide, bx, y, bw, h, SAND, line=CARD_LINE)
        rect(slide, bx, y, bw, 0.07, accent)
        text(slide, bx + 0.16, y + 0.2, bw - 0.32, 0.3, str(i + 1), size=20, color=accent, bold=True)
        text(slide, bx + 0.16, y + 0.62, bw - 0.32, h - 0.75,
             [[(title, {"bold": True, "color": INK, "size": title_size})], [(body, {"color": BODY, "size": body_size})]], line=1.13)
        if i < n - 1:
            arrow = rect(slide, bx + bw + 0.05, y + h / 2 - 0.13, gap - 0.1, 0.26, MUTED, shape=MSO_SHAPE.RIGHT_ARROW)
            arrow.line.fill.background()


def versus(slide, y, h, left, right, left_items, right_items, size=12.5):
    """Two columns: without (red) and with (green)."""
    for i, (title, items, accent) in enumerate(((left, left_items, RED), (right, right_items, GREEN))):
        x = 0.72 + i * 6.05
        rect(slide, x, y, 5.85, h, SAND, line=CARD_LINE)
        rect(slide, x, y, 5.85, 0.07, accent)
        text(slide, x + 0.22, y + 0.2, 5.4, 0.4, title, size=15, color=accent, bold=True)
        paras = [[("✕  " if i == 0 else "✓  ", {"bold": True, "color": accent, "size": size}), (a + " ", {"bold": True, "color": INK, "size": size}),
                  (b, {"color": BODY, "size": size})] for a, b in items]
        box = text(slide, x + 0.22, y + 0.7, 5.4, h - 0.85, paras, line=1.12)
        for p in box.text_frame.paragraphs:
            p.space_after = Pt(8)


def flow_runs() -> list[dict]:
    """The published evidence flow: one summary line per run (evidence/flow.jsonl)."""
    path = EVIDENCE / "flow.jsonl"
    if not path.exists():
        return []
    rows = []
    for line in path.read_text().splitlines():
        try:
            d = json.loads(line)
        except ValueError:
            continue
        if "label" in d:
            rows.append(d)
    return rows


def run_metrics(key: str) -> dict:
    """Metrics of the published run for one patient, from its event log."""
    path = EVIDENCE / f"run-{key}.jsonl"
    m = {"seconds": "–", "tool_calls": "–", "sql": "–", "searches": "–", "delegations": "–", "citations": "–", "verified": "–"}
    if not path.exists():
        return m
    for line in path.read_text().splitlines():
        try:
            e = json.loads(line)
        except ValueError:
            continue
        if e.get("type") == "done":
            m.update({k: e[k] for k in ("seconds", "tool_calls", "sql", "searches", "delegations")})
        if e.get("type") == "verify" and e.get("final", True):
            m["citations"] = e.get("citations", "–")
            m["verified"] = "passed" if e.get("passed") else "failed"
    return m


# ------------------------------------------------------------------- slides --

def build(template: Path, out: Path) -> None:
    prs = Presentation(str(template))
    drop_all_slides(prs)
    blank = layout(prs, "Light - Blank")

    # 1 · cover
    s = prs.slides.add_slide(layout(prs, "Title/Cover_5"))
    for ph in s.placeholders:
        idx = ph.placeholder_format.idx
        if idx == 0:
            ph.text = "Deep Agents on Your Own Data"
        elif ph.placeholder_format.type is not None and ph.name.startswith("Text Placeholder 2"):
            ph.text = "LangChain Deep Agents on Oracle AI Database"
        elif ph.name.startswith("Text Placeholder 3"):
            ph.text = "Federico Kamelhar"
        elif ph.name.startswith("Text Placeholder 4"):
            ph.text = "Senior Principal Architect, Agentic AI\nOracle\nOctober 2026"
    for ph in list(s.placeholders):
        if ph.has_text_frame and not ph.text_frame.text.strip():
            ph._element.getparent().remove(ph._element)
    rect(s, 2.62, 1.1, 0.02, 0.4, RGBColor(0x8F, 0xA6, 0xAD))
    s.shapes.add_picture(str(LOGOS / "langchain-white.png"), Inches(2.82), Inches(1.12), height=Inches(0.36))
    rect(s, 4.67, 1.1, 0.02, 0.4, RGBColor(0x8F, 0xA6, 0xAD))
    s.shapes.add_picture(str(LOGOS / "nvidia-white-mono.png"), Inches(4.87), Inches(1.12), height=Inches(0.36))
    notes(s, "Deep Agents on your own data: a LangChain Deep Agent, built with the open-source langchain-oracle packages, "
             "writes a pre-visit clinical brief from Oracle AI Database 26ai, inside an NVIDIA OpenShell sandbox. "
             "Every patient is synthetic.")

    # 1b · agenda
    s = prs.slides.add_slide(blank)
    header(s, "One story, one live demo", "Agenda")
    agenda = [
        ("Deep Agents", "what they are, why they fit enterprise processes", "3–4"),
        ("The method", "reverse-engineer a process: a clinic's pre-visit prep", "5–7"),
        ("The agent", "architecture, built with langchain-oracle", "8–10"),
        ("The data", "Oracle AI Database and how agents read it", "11–13"),
        ("Actions and policy", "propose, refuse, escalate, a doctor decides", "14–16"),
        ("Safety", "four layers; why OpenShell", "17–18"),
        ("Live demo", "switch to the browser, then back to the slides", "19–20"),
        ("What you saw, and proof", "the run in detail, memory, the gate, results, Codex", "21–29"),
    ]
    for i, (t, sub, pages) in enumerate(agenda):
        y = 1.8 + i * 0.6
        rect(s, 0.72, y, 11.9, 0.52, SAND if i % 2 == 0 else WHITE, line=CARD_LINE)
        rect(s, 0.72, y, 0.07, 0.52, RED if t in ("Deep Agents", "The method", "Live demo") else TEAL)
        text(s, 1.0, y + 0.13, 0.5, 0.35, f"{i + 1:02d}", size=15, color=RED, bold=True, font=MONO)
        text(s, 1.65, y + 0.12, 3.6, 0.35, t, size=16, color=INK, bold=True)
        text(s, 5.3, y + 0.15, 6.0, 0.35, sub, size=13.5, color=BODY)
        text(s, 11.3, y + 0.15, 1.2, 0.35, pages, size=12, color=MUTED, align=PP_ALIGN.RIGHT)
    notes(s, "The thread: a Deep Agent does real research and real work on a chart, and the database, the sandbox and a gate decide what it may do. "
             "The demo runs live in the middle; everything after it is the proof.")

    # 2b · what is a deep agent
    s = prs.slides.add_slide(blank)
    header(s, "LangChain Deep Agents", "What is a Deep Agent?")
    text(s, 0.72, 1.7, 11.9, 1.05, [[("An agent that runs long, multi-step operations on its own, from goal to finished result. ", {"bold": True, "color": INK, "size": 17}),
        ("Give it an objective, not a prompt: it works for minutes or hours, across dozens of steps and tools, and comes back with the work done "
         "and a record of how, stopping only where a person must decide.", {"color": BODY, "size": 14.5})]], line=1.15)
    traits = [
        ("It plans", "Turns the objective into a plan of steps, revises it as it learns, and stays on track from the first step to the last.", TEAL),
        ("It delegates", "Hands each part to a specialist with a narrow focus, the way a lead hands work to a team, in parallel where it can.", VIOLET),
        ("It works on your data", "Connects to the systems of record and the sources your people use, and shows where every finding came from.", OCHRE),
        ("It stays under control", "People approve the steps that matter; every step can be watched, replayed and audited.", RED),
    ]
    for i, (t, b, c) in enumerate(traits):
        x = 0.72 + i * 3.02
        rect(s, x, 3.05, 2.84, 2.0, SAND, line=CARD_LINE)
        rect(s, x, 3.05, 2.84, 0.07, c)
        text(s, x + 0.2, 3.27, 2.46, 2.0, [[(t, {"bold": True, "size": 16, "color": c})], [(b, {"size": 12.5, "color": BODY})]], line=1.18)
        for para in s.shapes[-1].text_frame.paragraphs:
            para.space_after = Pt(6)
    rect(s, 0.72, 5.35, 11.9, 0.95, TEAL)
    text(s, 0.98, 5.47, 11.4, 0.8, [[("A proven pattern, now open source. ", {"bold": True, "color": WHITE, "size": 13.5}),
        ("It is the design behind coding agents and deep-research products. LangChain packages it as \u201cthe batteries-included agent harness\u201d for long, "
         "multi-step work; langchain-oracle connects it to Oracle AI Database.", {"color": WHITE, "size": 13})]], line=1.15)
    notes(s, "Strategic framing, from LangChain's own positioning: Deep Agents is an agent harness for complex, multi-step, long-running work. "
             "Their argument is that the model loop is the same as any agent's; the difference is the architecture around it: a plan, specialists, "
             "a place to keep working notes, and a detailed brief for each role. Claude Code, Deep Research and Manus follow this pattern. "
             "Sources: docs.langchain.com/oss/python/deepagents, langchain.com/blog/deep-agents, github.com/langchain-ai/deepagents.")

    # 2c · why enterprises need them
    s = prs.slides.add_slide(blank)
    header(s, "Enterprise work is processes, not questions", "Why Deep Agents fit the enterprise")
    needs = [
        ("Many steps", "A real task is gather, analyse, research, decide, act, record. A Deep Agent plans it and works it through.", TEAL),
        ("Many systems", "Records, documents, reference sources. Specialists each own one source, with only the tools for it.", VIOLET),
        ("Many roles", "People with different powers. Agents mirror the roles, and the database enforces who may do what.", OCHRE),
        ("Sign-off and audit", "Nothing consequential happens without a person, and every step leaves a trail you can show a regulator.", RED),
    ]
    for i, (t, b, c) in enumerate(needs):
        x = 0.72 + (i % 2) * 6.05
        y = 1.9 + (i // 2) * 1.75
        rect(s, x, y, 5.85, 1.55, SAND, line=CARD_LINE)
        rect(s, x, y, 0.07, 1.55, c)
        text(s, x + 0.3, y + 0.2, 5.35, 1.3, [[(t, {"bold": True, "size": 16, "color": INK})], [(b, {"size": 12.5, "color": BODY})]], line=1.15)
    text(s, 0.72, 5.5, 11.9, 1.05, [[("Why this makes reverse-engineering work: ", {"bold": True, "color": RED, "size": 14}),
        ("a Deep Agent’s parts map one-to-one onto an existing process. The plan is the process’s steps, the sub-agents are its roles, "
         "the tools are its systems, the approvals are its sign-offs. You copy the process; you do not invent a new one.", {"color": INK, "size": 13.5})]], line=1.15)
    notes(s, "Chat assistants help a person answer a question. Deep Agents take on the process around the question. "
             "The blockers are not model quality; they are data access, separation of duties, approvals and audit. The rest of the talk is how each one is handled.")

    # 2d · reverse engineering a process
    s = prs.slides.add_slide(blank)
    header(s, "Start from how the work is done today, not from the model", "Reverse-engineer the process")
    steps(s, 0.72, 1.9, 11.9, 2.75, [
        ("Observe", "Follow the process as people run it: who does each step, with which data, under which rule.", TEAL),
        ("Decompose", "Split it into roles and steps. Each role becomes a specialist agent; the coordinator becomes the lead.", VIOLET),
        ("Map the controls", "Each data source → a tool with one store. Each rule → a database policy. Each sign-off → a human approval.", OCHRE),
        ("Prove it", "Run it on synthetic cases. A gate checks every output; every step leaves evidence a reviewer can replay.", GREEN),
    ])
    text(s, 0.72, 4.95, 11.9, 1.5, [[("Why this order matters. ", {"bold": True, "color": INK, "size": 13.5}),
        ("The process already encodes decades of judgement about who may do what. Copying it gives the agents a structure people trust, "
         "and it tells you where the controls go: wherever a person signs today, the agent stops and asks.", {"color": BODY, "size": 13})]], line=1.15)
    rect(s, 0.72, 5.9, 11.9, 0.62, TEAL)
    text(s, 0.98, 6.02, 11.4, 0.45, [[("The same mapping works beyond healthcare: ", {"bold": True, "color": WHITE, "size": 12.5}),
        ("claims handling, procurement approvals, KYC reviews, incident post-mortems: any process with roles, sources, rules and sign-offs.", {"color": WHITE, "size": 12.5})]])
    notes(s, "This is the method behind the demo. We did not start with prompts; we started with how a clinic prepares for a visit, "
             "and turned each role into an agent and each rule into a database policy. The next slides show that mapping.")

    # 3b · the process, reverse-engineered
    s = prs.slides.add_slide(blank)
    header(s, "Pre-visit preparation in a clinic, role by role", "The process, reverse-engineered")
    table(s, 0.72, 1.85, 11.9, [
        ["Step today", "Done today by", "In the demo", "Control"],
        ["Pull the chart: labs, meds, notes, trends", "Nurse or medical assistant", "chart-analyst (SQL + note search)", "sees one patient (row-level security)"],
        ["Check the guidelines", "Clinician", "guideline-researcher (NIH reference)", "cites only what it retrieved"],
        ["Check the evidence", "Clinician", "evidence-researcher (PubMed)", "cites only what it retrieved"],
        ["Order labs, message the patient, book a visit", "Care coordinator", "care-coordinator: proposes", "clinician approves (CP-01)"],
        ["Review a medication change", "Pharmacist", "medication-safety: the only agent allowed", "policy CP-03, in the database"],
        ["Sign the medication order", "Physician", "the doctor, in the web app", "policy CP-02: physician only"],
        ["Write the pre-visit summary", "Clinician", "lead: the cited brief", "gate checks every source id"],
    ], col_w=[3.5, 2.3, 3.3, 2.8], size=11.5)
    text(s, 0.72, 5.65, 11.9, 0.8, [[("Nothing new was invented for the agents: ", {"bold": True, "color": INK, "size": 13}),
        ("each one does a step a person does today, with that person’s data and that person’s limits. The people who sign today still sign.", {"color": BODY, "size": 13})]], line=1.15)
    repo(s, "agent/brief_agent.py (one prompt per role) · migrations/versions/0006-0008 (the controls)")
    notes(s, "Walk the table top to bottom: this is a real clinic workflow. The point for any enterprise process (claims, procurement, KYC, incident review) "
             "is the same mapping: roles to agents, sources to tools, rules to database policies, signatures to approvals.")

    # 3 · the scenario
    s = prs.slides.add_slide(blank)
    header(s, "Three synthetic patients, one question each", "The scenario: a pre-visit clinical brief")
    pats = [
        ("Patient X · SYN-X", "Type 2 diabetes with declining kidney function", "eGFR 68 → 47 in 18 months, UACR 22 → 148, daily ibuprofen, glipizide lows, retinal screening open since 2025.", TEAL),
        ("Patient Y · SYN-Y", "Heart failure, readmitted twice in 90 days", "K 5.6 on spironolactone + lisinopril + potassium chloride, KCl back on the discharge list, skipped diuretic doses, cost barriers.", RED),
        ("Patient Z · SYN-Z", "COPD with polypharmacy and a recent fall", "Tiotropium + ipratropium, oxybutynin, zolpidem + diphenhydramine, orthostatic drop, pulmonary rehab declined.", VIOLET),
    ]
    for i, (t, sit, body, c) in enumerate(pats):
        x = 0.72 + i * 4.12
        rect(s, x, 2.0, 3.86, 3.65, SAND, line=CARD_LINE)
        rect(s, x, 2.0, 3.86, 0.07, c)
        text(s, x + 0.22, 4.15, 3.4, 1.4, [[("The clinician asks: ", {"bold": True, "size": 11.5, "color": INK}),
                                           ("“" + QUESTIONS[i] + "”", {"italic": True, "size": 11.5, "color": BODY})]], line=1.15)
        text(s, x + 0.22, 2.22, 3.4, 3.3, [
            [(t, {"bold": True, "size": 16, "color": INK})],
            [(sit, {"bold": True, "size": 13, "color": c})],
            [(body, {"size": 12.5, "color": BODY})],
        ], line=1.18)
    rect(s, 0.72, 5.95, 11.9, 0.82, TEAL)
    text(s, 0.98, 6.07, 11.4, 0.6, [[("Every record is invented. ", {"bold": True, "color": WHITE, "size": 14}),
                                    ("The schema enforces it (CHECK is_synthetic = 'Y', ids start with SYN-), and the agent's input guard blocks real-looking identifiers. "
                                     "Reference text is public: MedQuAD (NIH) and PubMedQA.", {"color": WHITE, "size": 12.5})]])
    notes(s, "Each chart is hand-written to carry findings a careful clinician would raise. No real patient data is used anywhere, by design and by policy.")

    # 2 · why
    s = prs.slides.add_slide(blank)
    header(s, "Agents on data that matters", "The architecture")
    bullets(s, 0.72, 1.95, 6.55, 4.8, [
        ("The data is the hard part.", "Agents earn their keep on a company's own records, and those records are exactly what nobody wants an agent to wander through."),
        ("Deep Agents do real research.", "They plan, delegate to specialists, query and search, and write a document. That is more reach than a chat window, not less."),
        ("Safety has to be structural.", "The database decides which rows an agent sees; the sandbox decides which hosts it reaches; the key never enters it."),
        ("Grounding has to be checked.", "A brief is useful only when every claim points at a row, a note, or a source that exists."),
    ])
    figure(s, "figure1-architecture.png", 7.62, 1.92, 4.9,
           "Figure 1 — The Deep Agent runs inside an OpenShell sandbox; Oracle AI Database holds the chart, the vectors and the model.")
    notes(s, "Set up the problem: the value of agents is on your own data, and the risk is too. The answer in this talk is "
             "structural: the database, the sandbox and a grounding gate each refuse on their own.")

    # 4 · the deep agent
    s = prs.slides.add_slide(blank)
    header(s, "LangChain Deep Agents, built with langchain-oracle", "Plan, delegate, retrieve, write")
    figure(s, "figure2-deep-agent.png", 0.72, 1.9, 7.1)
    bullets(s, 8.15, 1.95, 4.5, 4.9, [
        ("Lead.", "Plans with write_todos, delegates with task, and returns a structured PreVisitBrief. It holds no data tools."),
        ("Chart analyst.", "Reads the relational chart with read-only SQL and the notes with hybrid search."),
        ("Two researchers.", "Each bound to one store, five or more searches, six to ten sources each."),
        ("Care coordinator.", "Proposes lab requests, a patient message and a follow-up; escalates medication concerns."),
        ("Medication safety.", "The only agent allowed to propose a medication change; a doctor approves it."),
        ("Models.", "GPT-5.5 leads; Gemini 2.5 Flash runs the specialists, both on OCI Generative AI."),
    ], size=13.5)
    notes(s, "create_deepagents_agent from langchain-oci composes the LangChain Deep Agents harness with Oracle datastores. "
             "Specialists each get a single-store tool set so routing is explicit.")

    # 5 · code
    s = prs.slides.add_slide(blank)
    header(s, "Thirty lines of langchain-oracle", "Build the agent")
    code(s, 0.72, 1.9, 7.4, 4.6, '''
from langchain_oci import create_deepagents_agent
from langchain_oci.datastores import ADB, create_datastore_tools
from langchain_oracledb.embeddings import OracleEmbeddings

# Embeddings run inside Oracle AI Database (ONNX model)
emb = OracleEmbeddings(conn=conn, params={
    "provider": "database", "model": "DA_OWNER.MINILM_L12"})

notes = ADB(dsn=dsn, user="DA_AGENT_Y", password=pw,
            table_name="PATIENT_NOTE_VEC")
analyst_tools = [*sql_tools, *create_datastore_tools(
    {"patient_notes": notes}, embedding_model=emb)]

agent = create_deepagents_agent(
    model=ChatOpenAI(model="openai.gpt-5.5",
                     base_url=OCI_OPENAI_ENDPOINT,
                     api_key=os.environ["OCI_GENAI_API_KEY"]),
    system_prompt=LEAD,
    subagents=[chart_analyst, guideline, evidence],
    middleware=[TodoListMiddleware(), *guards],
    response_format=ToolStrategy(PreVisitBrief),
)''', size=11)
    card(s, 8.45, 1.9, 4.17, 1.9, "One credential in the process",
         "OCI_GENAI_API_KEY is a placeholder inside the sandbox; OpenShell's proxy swaps in the real key on the way to OCI.", accent=GREEN)
    card(s, 8.45, 4.05, 4.17, 1.9, "Schema-shaped output",
         "ToolStrategy(PreVisitBrief) fixes the brief's sections; the runner renders and checks it before anyone reads it.", accent=TEAL)
    notes(s, "The agent is ordinary langchain-oracle code. The interesting parts are what it does not hold: no OCI key, no write grant.")

    # 6 · the data layer
    s = prs.slides.add_slide(blank)
    header(s, "One database holds everything the agents read", "Where the data lives")
    for i, (t, what, how, c) in enumerate([
        ("Tables: the chart", "Patients, conditions, medications, labs, vitals, visits, referrals, appointments, and clinical notes.",
         "The agents ask in SQL: “how did eGFR change?” becomes a SELECT with REGR_SLOPE.", TEAL),
        ("Vectors: searchable text", "Every clinical note, 147 NIH reference answers (MedQuAD) and 360 PubMed abstracts, chunked and embedded.",
         "The agents ask in words: “potassium and spironolactone” finds the passages that mean it.", VIOLET),
        ("The embedding model", "all-MiniLM-L12-v2 loaded into the database as an ONNX model.",
         "Text is turned into vectors where it lives. No note leaves the database to be embedded.", GREEN),
    ]):
        x = 0.72 + i * 4.12
        rect(s, x, 1.9, 3.86, 2.95, SAND, line=CARD_LINE)
        rect(s, x, 1.9, 3.86, 0.07, c)
        text(s, x + 0.22, 2.1, 3.42, 2.9, [[(t, {"bold": True, "size": 16, "color": INK})], [(what, {"size": 12.5, "color": BODY})],
                                          [("How agents use it: ", {"bold": True, "size": 12.5, "color": c}), (how, {"size": 12.5, "color": BODY})]], line=1.16)
        for para in s.shapes[-1].text_frame.paragraphs:
            para.space_after = Pt(8)
    rect(s, 0.72, 5.1, 11.9, 0.95, TEAL)
    text(s, 0.98, 5.22, 11.4, 0.8, [[("One login per patient. ", {"bold": True, "color": WHITE, "size": 14}),
        ("Patient Y's agents connect as DA_AGENT_Y. Whether they SELECT a table or search the vectors, the database itself "
         "returns only Y's rows. 300 background patients exist too, but agents only ever see them as percentiles (cohort_benchmark).", {"color": WHITE, "size": 12.5})]], line=1.15)
    text(s, 0.72, 6.28, 11.9, 0.25, "Schema, row-level security and policies are versioned Alembic migrations, so a reviewer can read exactly what the agents may do.", size=11, color=MUTED)
    repo(s, "migrations/versions/ · db/seed.py · db/embed.py · data/patients.py")
    notes(s, "Relational and vector side by side in Oracle AI Database 26ai. OracleAutonomousDatabaseLoader reads the note rows, OracleTextSplitter "
             "chunks them and OracleEmbeddings with provider=database embeds them in place; OracleVS writes the vector tables and Oracle Text adds keyword search.")

    # 6c · how agents read the database, and every tool
    s = prs.slides.add_slide(blank)
    header(s, "Six agents, each with only the tools its job needs", "How agents read the database")
    table(s, 0.72, 1.85, 11.9, [
        ["Agent", "Its tools", "What they do in the database"],
        ["Lead", "write_todos · task", "Nothing. It plans and delegates; every fact reaches the brief through a specialist."],
        ["Chart analyst", "describe_chart_tables · query_chart · search_patient_notes", "Reads the chart: any SELECT it writes (read-only, 200 rows), plus vector + keyword search of the notes."],
        ["Guideline researcher", "search_clinical_reference · get_document_…", "Searches the NIH reference vectors (MedQuAD)."],
        ["Evidence researcher", "search_research_evidence · get_document_…", "Searches the PubMed abstract vectors (PubMedQA)."],
        ["Care coordinator", "propose_lab_request · draft_patient_message · propose_follow_up · escalate_to_agent", "One fixed INSERT per proposal. Its propose_medication_change is refused by policy CP-03."],
        ["Medication safety", "list_escalations · propose_medication_change · decline_escalation · query_chart", "The only agent allowed to propose a medication change; re-checks the chart first."],
    ], col_w=[2.2, 4.6, 5.1], size=11)
    text(s, 0.72, 5.25, 11.9, 1.3, [[("Two paths, both guarded by the database. ", {"bold": True, "color": INK, "size": 13}),
        ("Reads: the model writes the SQL or the search, and row-level security filters it to one patient. "
         "Writes: the model fills in typed arguments, the tool binds them into one fixed INSERT, and triggers apply the policies. "
         "No agent can UPDATE, DELETE, approve or execute.", {"color": BODY, "size": 12.5})]], line=1.15)
    repo(s, "agent/sql_tools.py · agent/care_tools.py · agent/brief_agent.py (subagents=[…])")
    notes(s, "Search tools come from langchain-oci create_datastore_tools, one store per specialist so routing is explicit. "
             "The care tools stamp each session with the agent's name (CLIENT_IDENTIFIER) in tool code; the model cannot choose it.")

    # 7 · agent to SQL + VPD
    s = prs.slides.add_slide(blank)
    header(s, "The model can write any SELECT; the database decides the rows", "Agent to SQL, safely")
    steps(s, 0.72, 1.9, 11.9, 2.35, [
        ("The analyst writes SQL", "Any SELECT it decides on: trends, slopes, joins. Whether or not it adds a patient filter, the database applies one.", TEAL),
        ("Oracle rewrites it", "Row-level security (VPD) sees the login DA_AGENT_Y and appends WHERE patient_id = 'SYN-Y' to every table, before it runs.", RED),
        ("Only Y comes back", "The result is Y's labs. Ask for X's rows and you get zero, not an error the model could argue with.", GREEN),
    ])
    table(s, 0.72, 4.5, 7.1, [
        ["Tried as an agent login", "Database answer"],
        ["SELECT every patient", "1 row: its own"],
        ["Read another patient's labs", "0 rows"],
        ["UPDATE a medication", "ORA-41900 refused"],
        ["Insert a non-synthetic patient", "ORA-02290 refused"],
    ], col_w=[4.2, 2.9], size=11.5)
    text(s, 8.1, 4.55, 4.52, 2.0, [[("Why not filter in the tool? ", {"bold": True, "color": INK, "size": 13}),
        ("A tool filter is code the model's SQL could route around (a subquery, a view). The database applies the policy to every statement, "
         "whoever wrote it. scripts/verify.py --rls-only runs 25 such checks: 25 of 25 pass.", {"color": BODY, "size": 12.5})]], line=1.15)
    repo(s, "migrations/versions/0004_row_level_security.py · agent/sql_tools.py · scripts/verify.py --rls-only")
    notes(s, "DA_PATIENT_SCOPE returns 1=1 for DA_OWNER, DA_CLINICIAN and DA_PHYSICIAN, patient_id = 'SYN-<k>' for DA_AGENT_<k>, and 1=0 for anyone else. "
             "The same policy covers the PATIENT_NOTE_VEC vector table through JSON_VALUE(metadata, '$.patient_id').")

    # 7b · from brief to action
    s = prs.slides.add_slide(blank)
    header(s, "The agents draft the work; people decide; the database does it", "From brief to action")
    steps(s, 0.72, 1.9, 11.9, 2.6, [
        ("Agent proposes", "For Y: a BMP and BNP before the visit, a portal message to Jordan, a heart-failure visit within 7 days. Stored as “proposed”; the agent cannot change that.", OCHRE),
        ("Clinician reviews", "Each proposal is a card with its reason and citations. Edit the message, then approve or reject.", TEAL),
        ("Database executes", "DECIDE_CARE_ACTION creates the lab order, queues the message (simulated), requests the appointment.", GREEN),
        ("Memory records it", "The decision is written to Y's memory. The next brief builds on it and never re-proposes it.", VIOLET),
    ])
    rect(s, 0.72, 4.85, 11.9, 1.05, SAND, line=CARD_LINE)
    text(s, 0.98, 5.0, 11.4, 0.85, [[("Three logins, three powers. ", {"bold": True, "color": INK, "size": 13.5}),
        ("DA_AGENT_Y can only insert proposals for Y. DA_CLINICIAN can approve routine actions. DA_PHYSICIAN can also approve medication changes. "
         "Every step lands in care_action_event: proposed › approved › executed, with who did it.", {"color": BODY, "size": 12.5})]], line=1.15)
    repo(s, "agent/care_tools.py · migrations/versions/0006_care_actions.py · ui/server.py (/api/actions) · web/src/components/ActionsPanel.tsx")
    notes(s, "Nothing is sent anywhere real: the portal outbox is a table and every patient is synthetic. The runner's gate rejects a brief that comes without proposals, so every demo run has actions to decide.")

    # 7b2 · policy: an agent is refused and escalates
    s = prs.slides.add_slide(blank)
    header(s, "An agent is refused by policy, escalates to another agent, and a doctor decides", "Policy in action")
    steps(s, 0.72, 1.85, 11.9, 2.45, [
        ("Coordinator tries", "Y's potassium is 5.6 on spironolactone. The care coordinator proposes: hold potassium chloride.", OCHRE),
        ("CP-03 refuses", "The database: only the medication-safety agent may propose a medication change. ORA-20014, logged.", RED),
        ("It escalates", "escalate_to_agent opens an escalation for medication-safety, with the cited reason.", OCHRE),
        ("Med-safety reviews", "Re-reads the labs and the reference itself, then proposes the hold on that escalation.", TEAL),
        ("CP-02: doctor only", "Stored as needs_physician. The clinician's approval is refused (ORA-20012); the doctor's executes it.", GREEN),
    ], gap=0.22, title_size=13, body_size=11)
    text(s, 0.72, 4.6, 5.8, 1.7, [[("Why two agents? ", {"bold": True, "color": INK, "size": 13}),
        ("Separation of duties, as with a pharmacist's check: the agent that coordinates care is not the one that changes medications, "
         "and neither is the one that approves.", {"color": BODY, "size": 12.5})]], line=1.15)
    text(s, 6.82, 4.6, 5.8, 1.7, [[("Why in the database? ", {"bold": True, "color": INK, "size": 13}),
        ("A prompt can be ignored. The trigger reads which agent is calling (set by tool code, not the model) and refuses; "
         "the refusal is logged even though the insert rolls back.", {"color": BODY, "size": 12.5})]], line=1.15)
    rect(s, 0.72, 5.75, 11.9, 0.75, TEAL)
    text(s, 0.98, 5.88, 11.4, 0.55, [[("See it live: ", {"bold": True, "color": WHITE, "size": 13}),
        ("run a brief, open Actions. The chain sits above the cards, read from policy_event and agent_escalation; the sandbox console prints "
         "“✕ POLICY CP-03” the moment the database refuses.", {"color": WHITE, "size": 12})]], line=1.12)
    repo(s, "migrations/versions/0007_care_policy.py · 0008_agent_policy.py · agent/care_tools.py · web/src/components/PolicyChain.tsx")
    notes(s, "care_policy holds CP-01a/b/c (clinician), CP-02 (physician) and CP-03 (allowed_proposer = medication-safety). "
             "policy_event is written by an autonomous transaction, so the refusal survives the rollback. scripts/verify.py checks all of it; scripts/demo_e2e.py asserts it for X, Y and Z.")

    # 7b3 · the escalation, as it ran
    s = prs.slides.add_slide(blank)
    header(s, "Patient Y, recorded on a real run in the sandbox", "The escalation, live")
    figure(s, "../screenshots/09-policy-chain.png", 0.72, 1.75, 11.9)
    text(s, 0.72, 3.47, 11.9, 0.3, "In the web app: refused by CP-03 → escalation #7 → medication-safety proposes action #89 → the doctor approves; the database executes it.", size=10.5, color=BODY)
    figure(s, "../screenshots/10-console-policy.png", 0.72, 3.9, 11.9)
    text(s, 0.72, 5.03, 11.9, 0.3, "In the sandbox console, the same moment: the database's refusal, the escalation, and the network decision behind each step.", size=10.5, color=BODY)
    text(s, 0.72, 5.5, 11.9, 1.0, [[("Nobody scripted this. ", {"bold": True, "color": INK, "size": 13}),
        ("The coordinator chose the concern (spironolactone with a falling eGFR), the database refused it, and the medication-safety agent re-read the chart and the reference "
         "before proposing a hold. The clinician's approval was refused (ORA-20012); the doctor's approval executed it as a physician order.", {"color": BODY, "size": 12.5})]], line=1.15)
    repo(s, "evidence/runs/console-Y-first.log · docs/screenshots/04-actions.png · scripts/demo_e2e.py")
    notes(s, "This is the fallback if the live demo misbehaves: the same chain, captured from a real run. Every line in the console is from inside the sandbox.")

    # 8 · safety net
    s = prs.slides.add_slide(blank)
    header(s, "Measured from inside the sandbox", "Four layers, none asks the agent to behave")
    figure(s, "figure4-safety-net.png", 0.72, 1.85, 7.25)
    bullets(s, 8.3, 1.95, 4.35, 4.9, [
        ("Sandbox egress.", "Two surfaces granted; example.com, Object Storage and PyPI refused at connect; PUT gets 403."),
        ("Database.", "One patient per agent login; proposals only, under per-agent policy; non-synthetic rows rejected."),
        ("Input guard.", "Blocks SSN, phone, email, MRN, date of birth and other patients' ids."),
        ("Brief gate.", "Structure, length and every cited id, checked before acceptance."),
    ], size=13.5)
    notes(s, "Each layer has its own transcript: safety_probe.sh 7/7, verify.py --rls-only 25/25, guard checks 7/7.")

    # 8b · why openshell
    s = prs.slides.add_slide(blank)
    header(s, "The agent runs model-chosen code next to medical records", "Why OpenShell")
    versus(s, 1.85, 3.55, "On a laptop or a server, without a sandbox", "Inside an OpenShell sandbox",
        [("The real API key", "sits in the process; one leaked log or prompt injection and it is gone."),
         ("Any host is reachable.", "A note that says “send this chart to paste.site” is one HTTP call away."),
         ("Any package installs.", "pip can pull whatever the model asks for."),
         ("Nothing is recorded", "about where the agent connected."),
         ("State lingers:", "passwords and files stay on the machine after the run.")],
        [("The key is a placeholder;", "the OpenShell proxy swaps in the real one on the way to OCI GenAI."),
         ("Two destinations only:", "OCI GenAI and the database, and only from python3.12. Everything else is refused at connect."),
         ("No package index:", "PyPI is refused; the image is built before the run."),
         ("Every connection is logged", "with ALLOWED or DENIED and the rule; the web app shows it live."),
         ("One sandbox per run,", "deleted at the end; the database login file is deleted as soon as it is read.")])
    text(s, 0.72, 5.62, 11.9, 0.95, [[("The database limits what the agent can read and write. ", {"bold": True, "color": INK, "size": 13}),
        ("OpenShell limits where anything it read can go, and what it can steal. They cover different failures; neither asks the agent to behave. "
         "scripts/safety_probe.sh tests it from inside: 7 of 7.", {"color": BODY, "size": 12.5})]], line=1.15)
    repo(s, "sandbox/policy.template.yaml · sandbox/profile/oci-genai-python.yaml · scripts/sandbox.sh · scripts/safety_probe.sh")
    notes(s, "Threat model in one line: the agent reads untrusted text (notes, abstracts) and acts on it with tools. Prompt injection is a when, not an if. "
             "OpenShell makes the blast radius the two allowed endpoints, with no key to steal.")

    # 8c · live demo divider
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, TEAL)
    text(s, 0.72, 2.2, 11.9, 0.4, "LIVE DEMO · SWITCH TO THE BROWSER", size=13, color=YELLOW, bold=True, spacing=220)
    text(s, 0.72, 2.7, 11.9, 1.0, "Patient Y, end to end", size=40, color=WHITE, bold=True)
    beats = ["Pick Patient Y and press Enter: the sandbox starts, the console shows only two destinations allowed",
             "Watch the five specialists: SQL, searches, the coordinator's “✕ POLICY CP-03”, the escalation",
             "Read the brief: every claim cited, every id checked in the database",
             "Actions: approve the labs as clinician; the medication change refuses you; approve it as the doctor",
             "Memory, then Reset demo: back to a clean slate in seconds"]
    text(s, 0.72, 3.85, 11.9, 3.0, [[(f"{i + 1}   ", {"bold": True, "color": YELLOW, "size": 15}), (b, {"color": WHITE, "size": 15})] for i, b in enumerate(beats)], line=1.5)
    text(s, 0.72, 7.08, 11.9, 0.3, FOOTER, size=8, color=RGBColor(0xC8, 0xD6, 0xDB))
    notes(s, "Before the talk: web app on http://127.0.0.1:8765, Reset demo pressed, OpenShell gateway up, scripts/preflight.sh green. "
             "A Y brief takes four to six minutes; while it runs, talk through the console and the trace. "
             "Fallback if the network fails: open evidence/runs/ and the screenshots on the next slides.")

    # 8d · welcome back
    s = prs.slides.add_slide(blank)
    header(s, "Back to the slides", "What you just saw")
    table(s, 0.72, 1.85, 11.9, [
        ["In the demo", "What it shows", "Where it comes from"],
        ["The plan, then five specialists at work", "a Deep Agent: plan, delegate, retrieve, write", "LangChain Deep Agents + langchain-oracle"],
        ["SQL the model wrote itself, one patient's rows", "the database decides what an agent reads", "row-level security"],
        ["“✕ POLICY CP-03”, then an escalation", "an agent refused by policy hands off to the agent that may", "care_policy, agent_escalation"],
        ["The clinician refused, the doctor approves", "the people who sign today still sign", "CP-02, DECIDE_CARE_ACTION"],
        ["Only two destinations in the console", "where data can go, and no key to steal", "NVIDIA OpenShell"],
        ["Every claim with an id", "a brief is accepted only when its sources exist", "the grounding gate"],
    ], col_w=[4.0, 4.6, 3.3], size=12)
    text(s, 0.72, 5.3, 11.9, 1.0, [[("Next: ", {"bold": True, "color": RED, "size": 13.5}),
        ("the same run beat by beat, the screens in detail, the memory, the gate that caught a real mistake, and the results across all three patients.", {"color": INK, "size": 13.5})]], line=1.15)
    notes(s, "Land the demo: point at each row and name the slide it maps to. If the live run is still going, leave the browser on the console and come back to it after the results slide.")

    # 9 · one run
    s = prs.slides.add_slide(blank)
    header(s, "The web app and Codex call the same script", "One run, beat by beat")
    figure(s, "figure5-sequence.png", 0.72, 1.85, 7.25)
    bullets(s, 8.3, 1.95, 4.35, 4.9, [
        ("Create.", "Sandbox with the provider and the policy; the key stays in the gateway."),
        ("Upload.", "Agent code and a 0600 login file, deleted as soon as it is read."),
        ("Run.", "Plan, delegate, query, search, return the structured brief."),
        ("Watch.", "Every hop is a NET:OPEN or HTTP:POST line with ALLOWED and the policy that allowed it."),
    ], size=13.5)
    notes(s, "OpenShell 0.1.2 note: the first settings poll reloads policy about ten seconds in; the script waits for it.")

    # 10 · the web app
    s = prs.slides.add_slide(blank)
    header(s, "Pick a patient, ask, watch the agents work", "The application")
    shots = ROOT / "docs" / "screenshots"
    if (shots / "02-agents-at-work.png").exists():
        figure(s, "../screenshots/02-agents-at-work.png", 0.72, 1.85, 5.8)
        figure(s, "../screenshots/06-inspect-trace.png", 6.82, 1.85, 5.8)
        text(s, 0.72, 5.7, 11.9, 0.9, [[("Overview: ", {"bold": True, "color": INK, "size": 12.5}),
            ("a guided flow (patient, agents at work, brief, actions, memory); the live plan beside the sandbox console, where each agent's step sits next to the network decision it caused.  ", {"size": 12.5}),
            ("Agent trace: ", {"bold": True, "color": INK, "size": 12.5}),
            ("every reasoning step, delegation, tool call with full arguments and result, the gate's verdicts, filterable by agent; any panel goes full screen.", {"size": 12.5})]])
    notes(s, "The UI is a thin client: FastAPI runs scripts/sandbox.sh and streams the agent channel and the console channel over SSE.")

    # 10b · actions
    s = prs.slides.add_slide(blank)
    header(s, "The clinician decides", "Proposed actions, approved and executed")
    if (shots / "04-actions.png").exists():
        figure(s, "../screenshots/04-actions.png", 0.72, 1.85, 7.25)
    bullets(s, 8.3, 1.95, 4.35, 4.9, [
        ("Each card", "shows the rationale with its citations, the drafted patient message (editable), the tests, or the follow-up window."),
        ("Approve and execute", "calls DECIDE_CARE_ACTION as DA_CLINICIAN; the card moves proposed → approved → executed."),
        ("Policy on the card.", "Above the cards, the chain: refused by CP-03, escalated, proposed by medication-safety, waiting for the doctor (CP-02)."),
        ("Every decision", "is written to the patient's memory, so the next brief builds on it and never re-proposes it."),
    ], size=13.5)
    notes(s, "Nothing is sent anywhere real: the portal outbox is a table, and every patient is synthetic.")

    # 7c · state and memory
    s = prs.slides.add_slide(blank)
    header(s, "langgraph-oracledb in the agent's own schema", "The agent remembers")
    figure(s, "figure9-memory.png", 0.72, 1.85, 7.25)
    bullets(s, 8.3, 1.95, 4.35, 4.9, [
        ("Checkpoints.", "OracleSaver persists every step of every run; repairs resume from state."),
        ("Long-term memory.", "OracleStore with an IVF vector index over in-database embeddings, mounted at /memories/ by StoreBackend."),
        ("What it remembers.", "Each accepted brief and each clinician decision. The next brief says what changed and never re-proposes what was done or rejected."),
    ], size=13.5)
    notes(s, "The persistence pattern is the one documented in langchain-oracle's deepagents guide: checkpointer=OracleSaver, store=OracleStore, backend=StoreBackend.")

    # 11 · grounding gate
    s = prs.slides.add_slide(blank)
    header(s, "Caught on a real run", "A brief is accepted only when its sources exist")
    figure(s, "figure7-grounding-gate.png", 0.72, 1.85, 7.25)
    bullets(s, 8.3, 1.95, 4.35, 4.9, [
        ("What happened.", "The lead cited MEDQUAD-00001 for a heart-failure fact. No search had returned it."),
        ("What the gate did.", "Looked every id up as the patient's own user, and sent the draft back with the missing id named."),
        ("Why it matters.", "Fluent and wrong is the expensive failure. Here it never reaches the clinician."),
    ], size=13.5)
    notes(s, "The real heart-failure reference rows are MEDQUAD-036xx. The gate runs inside the sandbox; scripts/verify.py repeats it afterwards.")

    # 12 · results
    s = prs.slides.add_slide(blank)
    header(s, "Four runs inside the sandbox, through the web app", "Results")
    rows = [["Run", "Seconds", "Tool calls", "SQL", "Searches", "Citations", "Proposals", "CP-03", "Escalation", "Doctor", "Gate"]]
    for r in flow_runs():
        rows.append([f"Patient {r['patient']} · {r['label'].split(',')[0]}", round(r.get("seconds") or 0), r.get("tool_calls"), r.get("sql"),
                     r.get("searches"), r.get("citations"), r.get("actions", 0),
                     "refused" if r.get("refused") else "–", str(r.get("escalation", "–")).replace(" accepted", " ✓"),
                     DOCTOR.get(r.get("doctor"), r.get("doctor", "–")), "passed" if r.get("passed") else "failed"])
    table(s, 0.72, 1.95, 11.9, rows, col_w=[2.3, 1.0, 0.95, 0.65, 0.95, 0.95, 0.85, 1.0, 1.25, 1.1, 0.9], size=11)
    text(s, 0.72, 4.05, 11.9, 0.3, "scripts/demo_e2e.py: DEMO E2E OK (41/41). CP-03: the database refused the care coordinator; escalation: accepted by medication-safety; "
         "doctor: the clinician was refused, the doctor approved and the database executed it (the second brief's change is left for the doctor).", size=10, color=BODY)
    card(s, 0.72, 4.6, 3.86, 1.75, "Sandbox boundary", "safety_probe.sh: 7 of 7. Key is a placeholder; only GenAI and the database are reachable; 0 denied during runs.", accent=GREEN)
    card(s, 4.84, 4.6, 3.86, 1.75, "Database guarantees", "verify.py --rls-only: 25 of 25. One patient per user, proposals only for that patient, medication changes only from medication-safety and only approved by the doctor.", accent=RED)
    card(s, 8.96, 4.6, 3.66, 1.75, "Grounding", "Every cited id resolves in the database as the patient's own user, checked twice: in the sandbox and afterwards.", accent=TEAL)
    notes(s, "All numbers come from evidence/run-*.jsonl and the probe and verify transcripts.")

    # 13 · built by codex
    s = prs.slides.add_slide(blank)
    header(s, "Anyone can rebuild and rerun this demo by handing one file to a coding agent", "Why Codex")
    versus(s, 1.85, 3.0, "Reproducing a demo like this by hand", "Handing codex/PROMPT.md to Codex",
        [("About 25 steps:", "users, grants, ONNX model, 8 migrations, seed, embed, image, provider, sandbox, runs, checks."),
         ("Each one is a chance", "to skip a check, or paste a password into a terminal history."),
         ("“It worked for me”", "is the only proof.")],
        [("One prompt, one spec:", "Codex migrates, seeds, embeds, builds, probes, runs X, Y and Z, verifies, reports."),
         ("Codex never sees a secret:", "it stops twice, and you type the ADMIN password and the GenAI key yourself."),
         ("Every step has a pass condition,", "and Codex stops instead of working around a failed one.")])
    steps(s, 0.72, 5.05, 11.9, 1.5, [
        ("Run 1 stopped", "The database refused the laptop's new IP. Codex reported it and stopped; no workaround.", RED),
        ("Run 2 passed", "Full build and three briefs from a clean slate, every check green.", GREEN),
        ("Run 3 passed", "Through the web app: WEBAPP OK, 7 of 7.", GREEN),
        ("Run 4 passed", "Current code: probe 7/7, database 25/25, Y brief through CP-03 → escalation → doctor, web app 7/7.", GREEN),
    ], gap=0.22, title_size=12.5, body_size=10)
    repo(s, "codex/SPEC.md · codex/PROMPT.md · AGENTS.md (hard rules) · evidence/codex/run1-4")
    notes(s, "The point of Codex is reproducibility with the same safety posture as the demo: a written spec instead of a README, a coding agent that executes it "
             "under rules it cannot override (AGENTS.md: no secrets, synthetic data only, never decide a care action, never edit code to pass a check), "
             "and transcripts as proof. It is the same method as the OpenShell kill-switch post.")

    # 14 · field notes
    s = prs.slides.add_slide(blank)
    header(s, "What cost real time", "Field notes")
    notes_rows = [
        ("python-oracledb 26.x thin", "OracleEmbeddings fails building VECTOR_ARRAY_T (DPY-3013). Pin oracledb<4."),
        ("Gemini on the OpenAI endpoint", "Rejects exclusiveMinimum in tool schemas; tool results must be JSON objects."),
        ("langchain-oci 0.3.2 search", "Routes by meaning across all stores, with no store argument: give each specialist its own."),
        ("deepagents 0.7", "No to-do list in the base stack: add TodoListMiddleware for a visible plan."),
        ("Middleware and recursion", "Each middleware adds graph nodes; eleven PII guards meant six model turns. Merge them."),
        ("OpenShell 0.1.2", "The first settings poll closes open connections about 10 s in: wait for it. ADB TLS needs tcp + tls: skip."),
        ("Ids from memory", "Models cite plausible ids they never retrieved: give SQL rows their vector id and check every id."),
        ("Gemini and Literal[...]", "A one-value Literal becomes JSON-schema const, which Gemini rejects: map it to a one-value enum."),
        ("Examples in tool docs", "A specialist copied the example subject from a tool description into a real escalation. Describe the shape, not an instance."),
        ("Policies the model can see", "Told a tool would be refused, the model skipped it. Escalation now requires a refusal recorded by the database."),
        ("GPT-6 on API keys", "Chat Completions refuses tools with reasoning on; /responses returned 404 for our API key. GPT-5.5 leads."),
    ]
    table(s, 0.72, 1.85, 11.9, [["Where", "What we found"]] + [list(r) for r in notes_rows], col_w=[3.3, 8.6], size=11)
    notes(s, "Each of these is in the repository's codex/SPEC.md troubleshooting table or a code comment.")

    # 15 · conclusion
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, TEAL)
    text(s, 0.72, 0.7, 11.9, 0.35, "CONCLUSION", size=12, color=YELLOW, bold=True, spacing=200)
    text(s, 0.72, 1.15, 11.9, 1.7, [[("Deep Agents on your own data, safely", {"size": 34, "bold": True, "color": WHITE})],
                                   [("A LangChain Deep Agent plans, delegates and writes a grounded brief from Oracle AI Database; the database, the sandbox and the gate each "
                                     "decide what it can do, and none of them asks the agent to cooperate.", {"size": 15, "color": WHITE})]], line=1.15)
    closing = [("1  Run it", "github.com/fede-kamel/deep-agents-oracle-ai-database — hand codex/PROMPT.md to Codex."),
             ("2  Build on it", "langchain-oracle: create_deepagents_agent, ADB datastores, OracleEmbeddings, OracleVS, Oracle Text."),
             ("3  Keep it safe", "Row-level security for data, OpenShell for egress and keys, a gate for grounding.")]
    for i, (t, b) in enumerate(closing):
        x = 0.72 + i * 4.12
        rect(s, x, 3.85, 3.86, 1.9, RGBColor(0x03, 0x45, 0x5C), line=RGBColor(0x0B, 0x63, 0x80))
        rect(s, x, 3.85, 3.86, 0.07, YELLOW)
        text(s, x + 0.22, 4.05, 3.42, 2.0, [[(t, {"bold": True, "size": 16, "color": WHITE})], [(b, {"size": 12.5, "color": RGBColor(0xD8, 0xE6, 0xEA)})]], line=1.15)
    text(s, 0.72, 6.6, 11.9, 0.5, "Every claim on these slides has a transcript in the repository's evidence/ folder.", size=15, color=YELLOW, bold=True)
    text(s, 0.72, 7.08, 11.9, 0.3, FOOTER, size=8, color=RGBColor(0xC8, 0xD6, 0xDB))
    notes(s, "Close on the method: structural safety plus a grounding gate, reproducible from a spec.")

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print("wrote", out, len(prs.slides._sldIdLst), "slides")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", type=Path, required=True)
    ap.add_argument("--out", type=Path, default=ROOT / "deck" / "Deep-Agents-Oracle-AI-Database-FY26.pptx")
    a = ap.parse_args()
    build(a.template.expanduser(), a.out)
