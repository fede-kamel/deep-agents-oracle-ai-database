"""Render the agent's Markdown brief as a Word document.

Handles the subset the brief uses: headings, a block quote banner, bullet and
numbered lists, paragraphs with **bold**, and [citation] chips, which become
shaded monospace runs so they stand out in print.
"""

from __future__ import annotations

import re

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor

INLINE = re.compile(r"(\*\*[^*]+\*\*|\[[A-Z][A-Za-z0-9:_#,\- ]+\])")
ORACLE_RED = RGBColor(0xC7, 0x46, 0x34)


def _runs(paragraph, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("[") and part.endswith("]"):
            run = paragraph.add_run(part)
            run.font.name = "Courier New"
            run.font.size = Pt(8.5)
            run.font.highlight_color = WD_COLOR_INDEX.GRAY_25
        else:
            paragraph.add_run(part)


def _table(doc, rows: list[str]) -> None:
    cells = [[c.strip() for c in r.strip("|").split("|")] for r in rows
             if not re.fullmatch(r"\|?[\s:|-]+\|?", r)]
    if not cells:
        return
    table = doc.add_table(rows=len(cells), cols=max(len(r) for r in cells))
    table.style = "Light Grid Accent 1"
    for r, row in enumerate(cells):
        for c, value in enumerate(row):
            para = table.cell(r, c).paragraphs[0]
            if r == 0:
                para.add_run(value).bold = True
            else:
                _runs(para, value)
    doc.add_paragraph()


def markdown_to_docx(markdown: str, target, *, footer: str = "") -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        i += 1
        if not stripped:
            continue
        if stripped.startswith("|"):
            rows = [stripped]
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            _table(doc, rows)
            continue
        if stripped.startswith("# "):
            h = doc.add_heading(stripped[2:], level=0)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x31, 0x2D, 0x2A)
        elif stripped.startswith("## "):
            h = doc.add_heading(stripped[3:], level=1)
            for r in h.runs:
                r.font.color.rgb = ORACLE_RED
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            run = p.add_run(stripped.lstrip("> ").strip())
            run.italic = True
            run.font.color.rgb = ORACLE_RED
        elif re.match(r"^[-*] ", stripped):
            _runs(doc.add_paragraph(style="List Bullet"), stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", stripped))
        else:
            _runs(doc.add_paragraph(), stripped)
    if footer:
        section = doc.sections[0]
        p = section.footer.paragraphs[0]
        p.text = footer
        for r in p.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0x6F, 0x69, 0x64)
    doc.save(target)
